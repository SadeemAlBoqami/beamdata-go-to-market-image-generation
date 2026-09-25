"""Blind human evaluation controls for saved benchmark outputs."""

import csv
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock

import gradio as gr


REPO_ROOT = Path(__file__).resolve().parents[2]

EVALUATION_RESULTS = (
    REPO_ROOT
    / "evaluation"
    / "data"
    / "five_model_human_evaluation.csv"
)

MODEL_LABELS = ["A", "B", "C", "D", "E"]

RUBRIC = [
    ("overall_quality", "Overall quality"),
    ("prompt_adherence", "Prompt adherence"),
    ("text_accuracy", "Text accuracy"),
    ("composition", "Composition"),
    ("marketing_usefulness", "Marketing usefulness"),
    ("visual_appeal", "Visual appeal"),
]


def build_model_evaluation(label: str) -> list:
    """Create 1–5 rubric controls for one blind model."""

    controls = []

    with gr.Accordion(
        f"Rate Model {label}",
        open=True,
    ):
        for field, title in RUBRIC:

            # Text may not exist in every benchmark prompt.
            choices = (
                ["N/A", 1, 2, 3, 4, 5]
                if field == "text_accuracy"
                else [1, 2, 3, 4, 5]
            )

            control = gr.Radio(
                choices=choices,
                label=title,
            )

            controls.append(control)

    return controls


def save_evaluation(
    evaluator_id: str,
    prompt_id: str,
    records: list[dict],
    scores: list,
) -> str:
    """Save or update one evaluator's scores for one prompt."""

    evaluator_id = (evaluator_id or "").strip()

    if not evaluator_id:
        return "Enter an evaluator ID before saving."

    if len(records) != 5:
        return "Five benchmark model outputs are required."

    expected_scores = len(MODEL_LABELS) * len(RUBRIC)

    if len(scores) != expected_scores:
        return "Evaluation form is incomplete."

    # Require every score.
    if any(score is None for score in scores):
        return "Complete all ratings before saving."

    EVALUATION_RESULTS.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    fieldnames = [
        "evaluator_id",
        "prompt_id",
        "blind_label",
        "model",
        *[field for field, _ in RUBRIC],
        "timestamp",
    ]

    new_rows = []

    rubric_size = len(RUBRIC)

    for index, label in enumerate(MODEL_LABELS):

        start = index * rubric_size
        model_scores = scores[
            start:start + rubric_size
        ]

        row = {
            "evaluator_id": evaluator_id,
            "prompt_id": prompt_id,
            "blind_label": label,
            "model": records[index].get(
                "model",
                "Unknown",
            ),
            "timestamp": datetime.now(
                timezone.utc
            ).isoformat(),
        }

        for (
            (field, _),
            score,
        ) in zip(RUBRIC, model_scores):

            # Store N/A as blank so numerical analysis stays clean.
            row[field] = (
                ""
                if score == "N/A"
                else score
            )

        new_rows.append(row)

    # ---------------------------------------------------------
    # Update existing evaluation instead of creating duplicates.
    # ---------------------------------------------------------

    existing_rows = []

    if EVALUATION_RESULTS.is_file():

        with EVALUATION_RESULTS.open(
            newline="",
            encoding="utf-8",
        ) as source:

            existing_rows = list(
                csv.DictReader(source)
            )

    # Remove previous scores from same evaluator/prompt.
    existing_rows = [
        row
        for row in existing_rows
        if not (
            row.get("evaluator_id") == evaluator_id
            and row.get("prompt_id") == prompt_id
        )
    ]

    existing_rows.extend(new_rows)

    with EVALUATION_RESULTS.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as destination:

        writer = csv.DictWriter(
            destination,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(existing_rows)

    return (
        f"Evaluation saved for {prompt_id} "
        f"by {evaluator_id}."
    )


def load_evaluation_scores(
    evaluator_id: str,
    prompt_id: str,
) -> list:
    """Load previously saved scores for one evaluator and prompt."""

    empty_scores = [
        None
        for _ in range(
            len(MODEL_LABELS) * len(RUBRIC)
        )
    ]

    evaluator_id = (evaluator_id or "").strip()

    if not evaluator_id:
        return empty_scores

    if not EVALUATION_RESULTS.is_file():
        return empty_scores

    with EVALUATION_RESULTS.open(
        newline="",
        encoding="utf-8",
    ) as source:
        rows = list(csv.DictReader(source))

    saved_rows = {
        row["blind_label"]: row
        for row in rows
        if (
            row.get("evaluator_id") == evaluator_id
            and row.get("prompt_id") == prompt_id
        )
    }

    if not saved_rows:
        return empty_scores

    scores = []

    for label in MODEL_LABELS:
        row = saved_rows.get(label)

        if row is None:
            scores.extend(
                [None] * len(RUBRIC)
            )
            continue

        for field, _ in RUBRIC:
            value = row.get(field, "")

            if value == "":
                if field == "text_accuracy":
                    scores.append("N/A")
                else:
                    scores.append(None)
            else:
                scores.append(int(value))

    return scores

# Pairwise votes use a separate append-only file; legacy five-model scores
# above remain readable and untouched.
PAIRWISE_RESULTS = REPO_ROOT / "evaluation" / "data" / "pairwise_human_evaluation.csv"
_PAIR_LOCK = Lock()
PAIR_SCORE_FIELDS = [
    f"{field}_{side}" for side in ("a", "b") for field, _ in RUBRIC
]
PAIR_FIELDS = [
    "evaluator_id", "prompt_id", "prompt", "image_a_model", "image_a_path",
    "image_b_model", "image_b_path", "preference", *PAIR_SCORE_FIELDS,
    "timestamp",
]


def _image_reference(record: dict) -> str:
    """Store a portable repository-relative image reference where possible."""
    image_path = Path(record["image_path"])
    try:
        return str(image_path.relative_to(REPO_ROOT))
    except ValueError:
        return str(image_path)


def save_pairwise_evaluation(
    evaluator_id: str,
    prompt_id: str,
    prompt: str,
    image_a: dict,
    image_b: dict,
    preference: str,
    scores: list,
) -> str:
    """Append a pairwise decision without altering earlier evaluation data."""
    evaluator_id = (evaluator_id or "").strip()
    if not evaluator_id:
        return "Enter an evaluator ID before saving."
    if preference not in ("A", "B", "Tie"):
        return "Choose A, B, or Tie before saving."
    if not image_a.get("success") or not image_b.get("success"):
        return "Both saved images must be available before voting."
    if len(scores) != len(PAIR_SCORE_FIELDS):
        return "Detailed scoring data is invalid."
    for index, score in enumerate(scores):
        if score is None:
            continue
        field = PAIR_SCORE_FIELDS[index]
        if str(score) not in {"1", "2", "3", "4", "5"} and not (
            field.startswith("text_accuracy_") and score == "N/A"
        ):
            return "Detailed scoring data is invalid."

    row = {
        "evaluator_id": evaluator_id,
        "prompt_id": prompt_id,
        "prompt": prompt,
        "image_a_model": image_a["model"],
        "image_a_path": _image_reference(image_a),
        "image_b_model": image_b["model"],
        "image_b_path": _image_reference(image_b),
        "preference": preference,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }
    row.update({
        field: "" if score is None else score
        for field, score in zip(PAIR_SCORE_FIELDS, scores)
    })
    with _PAIR_LOCK:
        PAIRWISE_RESULTS.parent.mkdir(parents=True, exist_ok=True)
        needs_header = not PAIRWISE_RESULTS.is_file() or PAIRWISE_RESULTS.stat().st_size == 0
        with PAIRWISE_RESULTS.open("a", newline="", encoding="utf-8") as destination:
            writer = csv.DictWriter(destination, fieldnames=PAIR_FIELDS)
            if needs_header:
                writer.writeheader()
            writer.writerow(row)
    return f"Evaluation saved for {prompt_id}: {image_a['model']} vs {image_b['model']}."


def load_pairwise_evaluation(
    evaluator_id: str, prompt_id: str, image_a: dict, image_b: dict,
) -> list:
    """Restore this evaluator's latest vote and optional scores for a pair."""
    empty = [None] * (1 + len(PAIR_SCORE_FIELDS))
    if not (evaluator_id or "").strip() or not PAIRWISE_RESULTS.is_file():
        return empty
    a_path, b_path = _image_reference(image_a), _image_reference(image_b)
    saved = None
    with _PAIR_LOCK:
        with PAIRWISE_RESULTS.open(newline="", encoding="utf-8") as source:
            for row in csv.DictReader(source):
                if (
                    row.get("evaluator_id") == evaluator_id.strip()
                    and row.get("prompt_id") == prompt_id
                    and row.get("image_a_model") == image_a["model"]
                    and row.get("image_a_path") == a_path
                    and row.get("image_b_model") == image_b["model"]
                    and row.get("image_b_path") == b_path
                ):
                    saved = row
    if saved is None:
        return empty
    scores = [
        (int(value) if value.isdigit() else value) if value else None
        for field in PAIR_SCORE_FIELDS
        for value in [saved.get(field, "")]
    ]
    return [saved.get("preference") or None, *scores]


# ---------------------------------------------------------------------------
# Five-model blind ranking
# ---------------------------------------------------------------------------
# Stored separately from legacy criterion scores and pairwise votes so earlier
# evaluation data remains untouched.

RANKING_RESULTS = (
    REPO_ROOT
    / "evaluation"
    / "data"
    / "five_model_ranking_evaluation.csv"
)

RANKING_XLSX = (
    REPO_ROOT
    / "evaluation"
    / "data"
    / "five_model_ranking_evaluation.xlsx"
)

_RANKING_LOCK = Lock()

RANKING_FIELDS = [
    "evaluator_id",
    "prompt_id",
    "prompt",
    "blind_label",
    "model",
    "image_path",
    "rank",
    "points",
    "timestamp",
]


def _ranking_image_reference(record: dict) -> str:
    """Return a portable image reference for either service record format."""
    raw = record.get("image_path") or record.get("image") or ""

    if not raw:
        return ""

    if isinstance(raw, (str, Path)):
        path = Path(raw)
        try:
            return str(path.resolve().relative_to(REPO_ROOT.resolve()))
        except (ValueError, OSError):
            return str(path)

    return str(raw)


def _validate_ranks(ranks: list) -> tuple[bool, list[int] | None]:
    if len(ranks) != 5 or any(rank is None for rank in ranks):
        return False, None

    try:
        normalized = [int(rank) for rank in ranks]
    except (TypeError, ValueError):
        return False, None

    if sorted(normalized) != [1, 2, 3, 4, 5]:
        return False, None

    return True, normalized


def _write_ranking_workbook(rows: list[dict]) -> None:
    """Rebuild the Excel export from the canonical ranking CSV rows."""
    try:
        from openpyxl import Workbook
    except ImportError:
        return

    workbook = Workbook()

    raw_sheet = workbook.active
    raw_sheet.title = "Raw Rankings"
    raw_sheet.append(RANKING_FIELDS)
    for row in rows:
        raw_sheet.append([row.get(field, "") for field in RANKING_FIELDS])
    raw_sheet.freeze_panes = "A2"
    raw_sheet.auto_filter.ref = raw_sheet.dimensions

    model_sheet = workbook.create_sheet("Model Summary")
    model_sheet.append([
        "model",
        "rankings",
        "average_rank",
        "total_points",
        "average_points",
        "first_place_count",
        "last_place_count",
    ])

    model_groups = {}
    for row in rows:
        model_groups.setdefault(row.get("model", "Unknown"), []).append(row)

    for model in sorted(model_groups):
        group = model_groups[model]
        numeric_ranks = [int(row["rank"]) for row in group if row.get("rank")]
        numeric_points = [int(row["points"]) for row in group if row.get("points")]
        model_sheet.append([
            model,
            len(numeric_ranks),
            round(sum(numeric_ranks) / len(numeric_ranks), 3) if numeric_ranks else "",
            sum(numeric_points),
            round(sum(numeric_points) / len(numeric_points), 3) if numeric_points else "",
            sum(rank == 1 for rank in numeric_ranks),
            sum(rank == 5 for rank in numeric_ranks),
        ])
    model_sheet.freeze_panes = "A2"

    prompt_sheet = workbook.create_sheet("Prompt Summary")
    prompt_sheet.append([
        "prompt_id",
        "model",
        "evaluator_count",
        "average_rank",
        "total_points",
        "first_place_count",
    ])

    prompt_groups = {}
    for row in rows:
        key = (row.get("prompt_id", ""), row.get("model", "Unknown"))
        prompt_groups.setdefault(key, []).append(row)

    for prompt_id, model in sorted(prompt_groups):
        group = prompt_groups[(prompt_id, model)]
        numeric_ranks = [int(row["rank"]) for row in group if row.get("rank")]
        numeric_points = [int(row["points"]) for row in group if row.get("points")]
        prompt_sheet.append([
            prompt_id,
            model,
            len({row.get("evaluator_id", "") for row in group}),
            round(sum(numeric_ranks) / len(numeric_ranks), 3) if numeric_ranks else "",
            sum(numeric_points),
            sum(rank == 1 for rank in numeric_ranks),
        ])
    prompt_sheet.freeze_panes = "A2"

    evaluator_sheet = workbook.create_sheet("Evaluator Summary")
    evaluator_sheet.append([
        "evaluator_id",
        "completed_prompts",
        "ranking_rows",
    ])

    evaluator_groups = {}
    for row in rows:
        evaluator_groups.setdefault(row.get("evaluator_id", ""), []).append(row)

    for evaluator in sorted(evaluator_groups):
        group = evaluator_groups[evaluator]
        evaluator_sheet.append([
            evaluator,
            len({row.get("prompt_id", "") for row in group}),
            len(group),
        ])
    evaluator_sheet.freeze_panes = "A2"

    method_sheet = workbook.create_sheet("Methodology")
    methodology_rows = [
        ("Evaluation method", "Blind five-model ranking"),
        ("Unit of evaluation", "One common benchmark prompt with five saved model outputs"),
        ("Blinding", "Model identities and technical metadata are hidden from evaluators"),
        ("Randomization", "Image order is deterministic per evaluator and prompt"),
        ("Ranking scale", "1 = best, 5 = worst; each rank is used exactly once"),
        ("Preference points", "Rank 1=5 points, Rank 2=4, Rank 3=3, Rank 4=2, Rank 5=1"),
        ("Interpretation", "Points are relative preference points, not absolute image-quality scores"),
    ]
    for key, value in methodology_rows:
        method_sheet.append([key, value])

    for sheet in workbook.worksheets:
        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            sheet.column_dimensions[column_cells[0].column_letter].width = min(
                max(max_length + 2, 12),
                60,
            )

    RANKING_XLSX.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(RANKING_XLSX)


def save_ranking_evaluation(
    evaluator_id: str,
    prompt_id: str,
    prompt: str,
    records: list[dict],
    ranks: list,
) -> str:
    """Save or update one evaluator's five-model ranking for one prompt."""
    evaluator_id = (evaluator_id or "").strip()

    if not evaluator_id:
        return "Enter an evaluator ID before saving."

    if len(records) != 5:
        return "Exactly five benchmark model outputs are required."

    if not all(record.get("success") for record in records):
        return "All five saved benchmark images must be available before ranking."

    valid, normalized = _validate_ranks(ranks)
    if not valid or normalized is None:
        return "Use each rank 1, 2, 3, 4, and 5 exactly once."

    RANKING_RESULTS.parent.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now(timezone.utc).isoformat()
    new_rows = []

    for label, record, rank in zip(MODEL_LABELS, records, normalized):
        new_rows.append({
            "evaluator_id": evaluator_id,
            "prompt_id": prompt_id,
            "prompt": prompt,
            "blind_label": label,
            "model": record.get("model", "Unknown"),
            "image_path": _ranking_image_reference(record),
            "rank": rank,
            "points": 6 - rank,
            "timestamp": timestamp,
        })

    with _RANKING_LOCK:
        existing_rows = []

        if RANKING_RESULTS.is_file():
            with RANKING_RESULTS.open(
                newline="",
                encoding="utf-8",
            ) as source:
                existing_rows = list(csv.DictReader(source))

        existing_rows = [
            row
            for row in existing_rows
            if not (
                row.get("evaluator_id") == evaluator_id
                and row.get("prompt_id") == prompt_id
            )
        ]

        existing_rows.extend(new_rows)

        with RANKING_RESULTS.open(
            "w",
            newline="",
            encoding="utf-8",
        ) as destination:
            writer = csv.DictWriter(
                destination,
                fieldnames=RANKING_FIELDS,
            )
            writer.writeheader()
            writer.writerows(existing_rows)

        _write_ranking_workbook(existing_rows)

    excel_note = (
        f" Excel updated at {RANKING_XLSX.relative_to(REPO_ROOT)}."
        if RANKING_XLSX.is_file()
        else " CSV saved; install openpyxl to enable the Excel export."
    )

    return f"Ranking saved for {prompt_id} by {evaluator_id}.{excel_note}"


def load_ranking_evaluation(
    evaluator_id: str,
    prompt_id: str,
    records: list[dict],
) -> list:
    """Restore saved ranks in the current blind display order."""
    empty = [None] * len(records)
    evaluator_id = (evaluator_id or "").strip()

    if not evaluator_id or not RANKING_RESULTS.is_file():
        return empty

    with _RANKING_LOCK:
        with RANKING_RESULTS.open(
            newline="",
            encoding="utf-8",
        ) as source:
            rows = [
                row
                for row in csv.DictReader(source)
                if (
                    row.get("evaluator_id") == evaluator_id
                    and row.get("prompt_id") == prompt_id
                )
            ]

    if not rows:
        return empty

    by_model = {row.get("model", ""): row for row in rows}

    result = []
    for record in records:
        saved = by_model.get(record.get("model", ""))
        if not saved or not saved.get("rank"):
            result.append(None)
            continue
        try:
            result.append(int(saved["rank"]))
        except ValueError:
            result.append(None)

    return result

